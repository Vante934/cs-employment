"""Unit tests for knowledge base file parser."""

from pathlib import Path

import pytest

from nanobot.kb.parser import ParseResult, parse_file
from nanobot.kb.store import KnowledgeBaseStore
from nanobot.session.sqlite_store import SqliteSessionStore


@pytest.fixture
def temp_db(tmp_path: Path) -> SqliteSessionStore:
    """Create a temporary database."""
    store = SqliteSessionStore(tmp_path)
    store.initialize(tmp_path)
    return store


@pytest.fixture
def kb_store(temp_db: SqliteSessionStore) -> KnowledgeBaseStore:
    """Create a KnowledgeBaseStore instance."""
    return KnowledgeBaseStore(temp_db)


@pytest.fixture
def kb_files_dir(tmp_path: Path) -> Path:
    """Create KB files directory."""
    kb_dir = tmp_path / "kb" / "files"
    kb_dir.mkdir(parents=True)
    return kb_dir


def create_kb_file(kb_store: KnowledgeBaseStore, kb_files_dir: Path, filename: str, content: bytes) -> tuple[str, Path]:
    """Helper to create a KB file record and actual file."""
    import hashlib
    
    file_hash = hashlib.sha256(content).hexdigest()
    file_path = kb_files_dir / f"{file_hash}{Path(filename).suffix}"
    file_path.write_bytes(content)
    
    file_id = kb_store.create_file_record(
        filename=filename,
        file_type=Path(filename).suffix[1:],
        file_size=len(content),
        file_path=str(file_path),
        file_hash=file_hash,
        status="pending"
    )
    
    return file_id, file_path


def test_parse_txt_file(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing a TXT file."""
    # Mock get_runtime_subdir to return our test directory
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    content = "Hello, world!\nThis is a test.".encode("utf-8")
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "test.txt", content)
    
    result = parse_file(file_id, kb_store)
    
    assert isinstance(result, ParseResult)
    assert result.success is True
    assert result.text == "Hello, world!\nThis is a test."
    assert result.error is None
    assert result.word_count == len(result.text)


def test_parse_md_file(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing a Markdown file."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    content = "# Header\n\nSome **markdown** content.".encode("utf-8")
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "test.md", content)
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is True
    assert "Header" in result.text
    assert "markdown" in result.text


def test_parse_empty_file(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing an empty file."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "empty.txt", b"")
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is True
    assert result.text == ""
    assert result.word_count == 0


def test_parse_unsupported_format(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing an unsupported file format."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "test.xyz", b"unknown format")
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is False
    assert result.text is None
    assert "unsupported" in result.error.lower()


def test_parse_file_not_found(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing when file doesn't exist."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    # Create record but don't create actual file
    file_id = kb_store.create_file_record(
        filename="missing.txt",
        file_type="txt",
        file_size=100,
        file_path=str(kb_files_dir / "missing.txt"),
        file_hash="fakehash",
        status="pending"
    )
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is False
    assert "does not exist" in result.error.lower() or "not found" in result.error.lower()


def test_parse_nonexistent_record(kb_store: KnowledgeBaseStore):
    """Test parsing with non-existent file_id."""
    result = parse_file("nonexistent-id", kb_store)
    
    assert result.success is False
    assert "not found" in result.error.lower()


def test_parse_already_parsed_skipped(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test that already parsed files are skipped."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    content = "Test content".encode("utf-8")
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "test.txt", content)
    
    # Mark as already parsed
    kb_store.update_file_status(file_id, "parsed", word_count=100)
    
    result = parse_file(file_id, kb_store, force=False)
    
    assert result.success is True
    assert result.text is None  # Skipped
    assert result.metadata is not None
    assert result.metadata.get("skipped") is True


def test_parse_force_reparse(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test forced re-parsing."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    content = "Test content".encode("utf-8")
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "test.txt", content)
    
    # Mark as already parsed
    kb_store.update_file_status(file_id, "parsed", word_count=100)
    
    result = parse_file(file_id, kb_store, force=True)
    
    assert result.success is True
    assert result.text is not None
    assert result.text == "Test content"


def test_parse_gbk_encoding(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing GBK encoded text."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    content = "中文测试".encode("gbk")
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "gbk.txt", content)
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is True
    # extract_text() should detect GBK and decode correctly
    assert result.text is not None
    assert len(result.text) > 0


def test_parse_symlink_rejected(kb_store: KnowledgeBaseStore, kb_files_dir: Path, tmp_path: Path, monkeypatch):
    """Test that symlinks are rejected."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    # Create real file
    real_file = kb_files_dir / "real.txt"
    real_file.write_text("content")
    
    # Create symlink
    symlink = kb_files_dir / "link.txt"
    try:
        symlink.symlink_to(real_file)
    except OSError:
        pytest.skip("Symlinks not supported on this platform")
    
    # Create record pointing to symlink
    file_id = kb_store.create_file_record(
        filename="link.txt",
        file_type="txt",
        file_size=7,
        file_path=str(symlink),
        file_hash="linkhash",
        status="pending"
    )
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is False
    assert "symlink" in result.error.lower()


def test_parse_path_outside_kb_dir(kb_store: KnowledgeBaseStore, tmp_path: Path, monkeypatch):
    """Test that files outside KB directory are rejected."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: tmp_path / "kb")
    
    # Create file outside KB directory
    outside_file = tmp_path / "outside.txt"
    outside_file.write_text("content")
    
    file_id = kb_store.create_file_record(
        filename="outside.txt",
        file_type="txt",
        file_size=7,
        file_path=str(outside_file),
        file_hash="outsidehash",
        status="pending"
    )
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is False
    assert "outside" in result.error.lower() or "directory" in result.error.lower()


def test_parse_directory_rejected(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test that directories are rejected."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    # Create directory
    test_dir = kb_files_dir / "testdir"
    test_dir.mkdir()
    
    file_id = kb_store.create_file_record(
        filename="testdir",
        file_type="txt",
        file_size=0,
        file_path=str(test_dir),
        file_hash="dirhash",
        status="pending"
    )
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is False
    assert "not a regular file" in result.error.lower()


def test_parse_pdf_file(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing a real PDF file."""
    pytest.importorskip("pypdf")
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    from pypdf import PdfWriter
    
    # Create a simple PDF
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    
    pdf_path = kb_files_dir / "test_temp.pdf"
    with open(pdf_path, "wb") as f:
        writer.write(f)
    
    import hashlib
    content = pdf_path.read_bytes()
    file_hash = hashlib.sha256(content).hexdigest()
    final_path = kb_files_dir / f"{file_hash}.pdf"
    pdf_path.rename(final_path)
    
    file_id = kb_store.create_file_record(
        filename="test.pdf",
        file_type="pdf",
        file_size=len(content),
        file_path=str(final_path),
        file_hash=file_hash,
        status="pending"
    )
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is True
    assert result.text is not None
    assert result.page_count == 1


def test_parse_corrupted_pdf(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing a corrupted PDF file."""
    pytest.importorskip("pypdf")
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    # Create fake PDF with invalid content
    content = b"%PDF-1.4\nInvalid PDF content"
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "bad.pdf", content)
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is False
    assert result.error is not None
    assert "[error:" in result.error.lower()


def test_parse_docx_file(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing a real DOCX file."""
    pytest.importorskip("docx")
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    from docx import Document
    
    # Create a simple DOCX
    doc = Document()
    doc.add_paragraph("Test paragraph 1")
    doc.add_paragraph("Test paragraph 2")
    
    docx_path = kb_files_dir / "test_temp.docx"
    doc.save(str(docx_path))
    
    import hashlib
    content = docx_path.read_bytes()
    file_hash = hashlib.sha256(content).hexdigest()
    final_path = kb_files_dir / f"{file_hash}.docx"
    docx_path.rename(final_path)
    
    file_id = kb_store.create_file_record(
        filename="test.docx",
        file_type="docx",
        file_size=len(content),
        file_path=str(final_path),
        file_hash=file_hash,
        status="pending"
    )
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is True
    assert result.text is not None
    assert "Test paragraph 1" in result.text
    assert "Test paragraph 2" in result.text


def test_parse_corrupted_docx(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing a corrupted DOCX file."""
    pytest.importorskip("docx")
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    # Create fake DOCX with invalid ZIP content
    content = b"PK\x03\x04Invalid DOCX"
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "bad.docx", content)
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is False
    assert result.error is not None
    assert "[error:" in result.error.lower()


def test_parse_xlsx_file(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing a real XLSX file."""
    pytest.importorskip("openpyxl")
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    from openpyxl import Workbook
    
    # Create a simple XLSX
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Name"
    ws["B1"] = "Value"
    ws["A2"] = "Test"
    ws["B2"] = 123
    
    xlsx_path = kb_files_dir / "test_temp.xlsx"
    wb.save(str(xlsx_path))
    
    import hashlib
    content = xlsx_path.read_bytes()
    file_hash = hashlib.sha256(content).hexdigest()
    final_path = kb_files_dir / f"{file_hash}.xlsx"
    xlsx_path.rename(final_path)
    
    file_id = kb_store.create_file_record(
        filename="test.xlsx",
        file_type="xlsx",
        file_size=len(content),
        file_path=str(final_path),
        file_hash=file_hash,
        status="pending"
    )
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is True
    assert result.text is not None
    assert "Name" in result.text
    assert "Value" in result.text


def test_parse_corrupted_xlsx(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test parsing a corrupted XLSX file."""
    pytest.importorskip("openpyxl")
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    # Create fake XLSX with invalid ZIP content
    content = b"PK\x03\x04Invalid XLSX"
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "bad.xlsx", content)
    
    result = parse_file(file_id, kb_store)
    
    assert result.success is False
    assert result.error is not None
    assert "[error:" in result.error.lower()


def test_parse_large_text_truncation(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test that text over 200KB is handled per document.py limits."""
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    # Create text larger than 200,000 characters
    large_text = "x" * 250_000
    content = large_text.encode("utf-8")
    file_id, _ = create_kb_file(kb_store, kb_files_dir, "large.txt", content)
    
    result = parse_file(file_id, kb_store)
    
    # Should succeed but text is truncated
    assert result.success is True
    assert result.text is not None
    assert len(result.text) < 250_000
    # document.py adds truncation message
    assert "truncated" in result.text.lower() or len(result.text) <= 200_000


def test_parse_wrong_extension(kb_store: KnowledgeBaseStore, kb_files_dir: Path, monkeypatch):
    """Test file with wrong extension (PDF content with .txt extension)."""
    pytest.importorskip("pypdf")
    monkeypatch.setattr("nanobot.kb.parser.get_runtime_subdir", lambda x: kb_files_dir.parent)
    
    from pypdf import PdfWriter
    
    # Create PDF content
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    
    pdf_path = kb_files_dir / "temp.pdf"
    with open(pdf_path, "wb") as f:
        writer.write(f)
    
    content = pdf_path.read_bytes()
    pdf_path.unlink()
    
    # Save with .txt extension
    import hashlib
    file_hash = hashlib.sha256(content).hexdigest()
    wrong_path = kb_files_dir / f"{file_hash}.txt"
    wrong_path.write_bytes(content)
    
    file_id = kb_store.create_file_record(
        filename="fake.txt",
        file_type="txt",
        file_size=len(content),
        file_path=str(wrong_path),
        file_hash=file_hash,
        status="pending"
    )
    
    result = parse_file(file_id, kb_store)
    
    # extract_text() tries to read as text, may fail or return garbage
    # Important: should not crash, result should be valid
    assert isinstance(result, ParseResult)
    # Either succeeds with binary garbage or fails gracefully
    if not result.success:
        assert result.error is not None
